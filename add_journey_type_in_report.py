
from __future__ import annotations

import re
import sys
from datetime import datetime
from pathlib import Path
from urllib.parse import unquote, urlparse

import pandas as pd
import requests
from openpyxl import load_workbook

SCRIPT_DIR = Path(__file__).resolve().parent
VIOLATION_LOOKUP_DIR = SCRIPT_DIR / "Violation Lookup"
sys.path.insert(0, str(VIOLATION_LOOKUP_DIR))

from violation_extractor import get_db_connection  # noqa: E402
from Violation_lookup import normalize_txn_id_for_comparison  # noqa: E402


# --- Configuration (edit before running) ---
PLAZA_NAME = "nathavalasa"
FROM_DATE = "2026-08-01"  # YYYY-MM-DD
TO_DATE = "2026-08-09"    # YYYY-MM-DD
INPUT_FOLDER = SCRIPT_DIR / "Nathavalasa"  # folder with Excel files to update
# EXISTING_ETC_FOLDER = SCRIPT_DIR / "chhapar_etc_files"   # Example: SCRIPT_DIR / "galia_etc_files"
EXISTING_ETC_FOLDER = None # Example: SCRIPT_DIR / "galia_etc_files"
# ETC file transaction ID column source: "exit", "agency", or "cch"
ETC_LOOKUP_ID_COLUMN_SOURCE = "exit"
# Scan first N rows in ETC files to locate the header row
HEADER_SCAN_ROW_COUNT = 10
# -------------------------------------------


JOURNEY_TYPE_COLUMN = "Journey Type"
NOT_FOUND_VALUE = "NA"

CCH_TRAN_ID_COLUMN_NAMES = [
    "CCH Tran Id",
    "CCH Tran ID",
    "CCH Tran Id ",
    "CCH_Tran_Id",
    "CCH TXN NO",
    "CCH TXN NO.",
    "CCH_TXN_NO",
    "CCH Transaction Id",
    "CCH Transaction ID",
    "CCH Transaction No",
    "CCH Transaction No.",
]

EXIT_TXN_ID_COLUMN_NAMES = [
    "Exit Txn Id",
    "Exit Txn ID",
    "Exit Txn Id ",
    "Exit_Txn_Id",
    "Exit Transaction Id",
    "Exit Transaction ID",
    "EXIT TXN ID",
]

AGENCY_TRAN_ID_COLUMN_NAMES = [
    "Agency Tran Id",
    "Agency Tran ID",
    "Agency Tran Id ",
    "Agency_Tran_Id",
    "Agency Txn Id",
    "Agency Txn ID",
    "Agency Txn Id ",
    "Agency_Txn_Id",
    "Agency Transaction Id",
    "Agency Transaction ID",
    "AGENCY TRAN ID",
    "AGENCY TXN ID",
    "Agency Txn Id",
    "AgencyTxnId",
    "AgencyTxnID",
]

ETC_LOOKUP_ID_COLUMN_OPTIONS = {
    "exit": EXIT_TXN_ID_COLUMN_NAMES,
    "agency": AGENCY_TRAN_ID_COLUMN_NAMES,
    "cch": CCH_TRAN_ID_COLUMN_NAMES,
}

MOP_COLUMN_NAMES = [
    "MOP",
    "MOP ",
    "Mop",
    "Mode of Payment",
    "Mode Of Payment",
    "Trip Type",
    "Trip Type ",
    "Trip Type\t",
    "TripType",
    "TripType ",
    "TripType\t",
    "Trip_Type",
    "Trip_Type ",
    "Trip_Type\t",
    "Fare Type",
    "FareType",
    "Fare_Type",
    "Journey Type",
    "JourneyType",
    "Journey_Type",
]

ETC_HEADER_HINT_COLUMN_NAMES = [
    "Sr. No.",
    "Toll File Name",
    "Plaza Id",
    "Plaza Name",
    "Lane Id",
    "Tag Id",
    "CCHTxn ID",
    "Reference ID",
    "Licence Plate No",
    "Tag VC",
    "AVC",
    "Mapper VC",
    "MVC Code",
    "Audit VC",
    "Toll Amount",
    "Transaction Amount",
    "Tag Read Date Time",
    "Transaction Date",
    "Processed Date",
    "Transaction Status",
    "Reason",
    "Fare Type",
    "Reconciled Date",
    "Reconcilation File Name",
    "Settled Amount",
    "Transaction Type",
    "Is Violation?",
]


def sanitize_folder_name(value: str) -> str:
    """Sanitize text for folder/file names."""
    safe = re.sub(r'[<>:"/\\|?*]', "_", str(value).strip())
    return safe.strip("._") or "plaza"


def get_etc_files_folder(plaza_name: str) -> Path:
    """Return download folder: {plaza_name}_etc_files."""
    folder = SCRIPT_DIR / f"{sanitize_folder_name(plaza_name)}_etc_files"
    folder.mkdir(parents=True, exist_ok=True)
    return folder


def get_etc_lookup_id_column_names() -> list[str]:
    """Return the ETC transaction ID column list selected in configuration."""
    source = str(ETC_LOOKUP_ID_COLUMN_SOURCE).strip().lower()
    if source not in ETC_LOOKUP_ID_COLUMN_OPTIONS:
        valid = ", ".join(sorted(ETC_LOOKUP_ID_COLUMN_OPTIONS))
        raise ValueError(
            f"Invalid ETC_LOOKUP_ID_COLUMN_SOURCE='{ETC_LOOKUP_ID_COLUMN_SOURCE}'. "
            f"Expected one of: {valid}"
        )
    return ETC_LOOKUP_ID_COLUMN_OPTIONS[source]


def find_column(df: pd.DataFrame, column_names: list[str]) -> str | None:
    """Find a column using normalized exact then partial name match."""
    normalized_columns = [
        (col, normalize_header_cell(col))
        for col in df.columns
    ]

    for col_name in column_names:
        alias = normalize_header_cell(col_name)
        if not alias:
            continue

        for col, normalized_col in normalized_columns:
            if normalized_col == alias:
                return col

        for col, normalized_col in normalized_columns:
            if alias in normalized_col or normalized_col in alias:
                return col
    return None


def normalize_header_cell(value) -> str:
    """Normalize a header cell for comparison."""
    if pd.isna(value):
        return ""
    text = str(value).replace("\n", " ").replace("\r", " ").replace("\t", " ")
    return " ".join(text.split()).strip().lower()


def cell_matches_column_names(cell_value, column_names: list[str]) -> bool:
    """Return True when a cell matches any known column alias."""
    cell_norm = normalize_header_cell(cell_value)
    if not cell_norm:
        return False

    for col_name in column_names:
        alias = normalize_header_cell(col_name)
        if not alias:
            continue
        if cell_norm == alias or alias in cell_norm or cell_norm in alias:
            return True
    return False


def row_has_required_columns(row: pd.Series, column_name_lists: list[list[str]]) -> bool:
    """Return True when one row contains at least one match for each column list."""
    for column_names in column_name_lists:
        if not any(cell_matches_column_names(value, column_names) for value in row):
            return False
    return True


def count_matching_column_groups(row_values, column_name_lists: list[list[str]]) -> int:
    """Count how many header alias groups match within a row."""
    return sum(
        1
        for column_names in column_name_lists
        if any(cell_matches_column_names(value, column_names) for value in row_values)
    )


def detect_excel_header_row_index(
    file_path: Path,
    txn_id_column_names: list[str],
    scan_rows: int = HEADER_SCAN_ROW_COUNT,
) -> int | None:
    """
    Detect the header row in an Excel sheet using openpyxl.
    This is more reliable for files with merged title rows and irregular structure.
    """
    try:
        workbook = load_workbook(file_path, read_only=True, data_only=True)
    except Exception as exc:
        print(f"    Warning: Could not scan Excel header rows in {file_path.name}: {exc}")
        return None

    try:
        worksheet = workbook[workbook.sheetnames[0]]
        candidate_lists = [
            txn_id_column_names,
            MOP_COLUMN_NAMES,
            ETC_HEADER_HINT_COLUMN_NAMES,
        ]

        best_row_idx: int | None = None
        best_score = -1

        for row_idx, row in enumerate(
            worksheet.iter_rows(
                min_row=1,
                max_row=scan_rows,
                values_only=True,
            )
        ):
            if not any(normalize_header_cell(value) for value in row):
                continue

            if not any(
                cell_matches_column_names(value, txn_id_column_names)
                for value in row
            ):
                continue

            score = count_matching_column_groups(row, candidate_lists)
            if score > best_score:
                best_score = score
                best_row_idx = row_idx

            if score == len(candidate_lists):
                return row_idx

        return best_row_idx
    finally:
        workbook.close()


def detect_header_row_index(
    file_path: Path,
    column_name_lists: list[list[str]],
    scan_rows: int = HEADER_SCAN_ROW_COUNT,
) -> int | None:
    """
    Scan the first rows of a file and return the row index to use as header.
    The header row must contain all required column groups.
    """
    suffix = file_path.suffix.lower()
    txn_id_column_names = column_name_lists[0] if column_name_lists else []

    if suffix in (".xlsx", ".xlsm", ".xls") and txn_id_column_names:
        header_row_idx = detect_excel_header_row_index(
            file_path,
            txn_id_column_names,
            scan_rows=scan_rows,
        )
        if header_row_idx is not None:
            return header_row_idx

    try:
        if suffix == ".csv":
            sample = pd.read_csv(
                file_path,
                header=None,
                nrows=scan_rows,
                dtype=str,
                keep_default_na=False,
            )
        elif suffix in (".xlsx", ".xlsm", ".xls"):
            sample = pd.read_excel(
                file_path,
                header=None,
                nrows=scan_rows,
                engine="openpyxl",
                dtype=str,
            )
        else:
            return None
    except Exception as exc:
        print(f"    Warning: Could not scan header rows in {file_path.name}: {exc}")
        return None

    for row_idx in range(min(scan_rows, len(sample))):
        if row_has_required_columns(sample.iloc[row_idx], column_name_lists):
            return row_idx

    return None


def read_tabular_file(
    file_path: Path,
    header: int | None = 0,
    nrows: int | None = None,
) -> pd.DataFrame | None:
    """Read an Excel or CSV file into a DataFrame."""
    suffix = file_path.suffix.lower()
    try:
        if suffix == ".csv":
            return pd.read_csv(
                file_path,
                dtype=str,
                keep_default_na=False,
                header=header,
                nrows=nrows,
            )
        if suffix in (".xlsx", ".xlsm", ".xls"):
            return pd.read_excel(
                file_path,
                engine="openpyxl",
                dtype=str,
                header=header,
                nrows=nrows,
            )
    except Exception as exc:
        print(f"    Warning: Could not read {file_path.name}: {exc}")
    return None


def read_etc_tabular_file(
    file_path: Path,
    txn_id_column_names: list[str],
) -> tuple[pd.DataFrame | None, int | None]:
    """
    Read an ETC file using a header row detected within the first scan rows.
    Returns the DataFrame and the detected header row index (0-based).
    """
    header_row_idx = detect_header_row_index(
        file_path,
        [txn_id_column_names],
        scan_rows=HEADER_SCAN_ROW_COUNT,
    )
    df = read_tabular_file(
        file_path,
        header=header_row_idx if header_row_idx is not None else 0,
    )
    return df, header_row_idx


def parse_date_yyyy_mm_dd(value: str) -> None:
    """Validate YYYY-MM-DD date string."""
    datetime.strptime(value.strip(), "%Y-%m-%d")


def download_filename_from_url(url: str, date_value, index: int) -> str:
    """Build a stable local filename for a downloaded ETC file."""
    parsed = urlparse(url)
    basename = Path(unquote(parsed.path)).name
    if not basename or basename in (".", "/"):
        basename = f"etc_file_{index}.xlsx"

    date_part = ""
    if date_value is not None and not pd.isna(date_value):
        if isinstance(date_value, (datetime, pd.Timestamp)):
            date_part = pd.Timestamp(date_value).strftime("%Y-%m-%d")
        else:
            date_part = str(date_value).strip()[:10]

    if date_part:
        stem = Path(basename).stem
        suffix = Path(basename).suffix or ".xlsx"
        basename = f"{date_part}_{stem}{suffix}"

    return sanitize_folder_name(basename).replace(" ", "_")


def download_etc_files(
    plaza_name: str,
    from_date: str,
    to_date: str,
    output_folder: Path,
    ) -> list[Path]:
    """Download ETC files from submissions for the plaza and date range."""
    conn = get_db_connection()
    if not conn:
        raise RuntimeError("Failed to connect to database.")

    downloaded_paths: list[Path] = []

    try:
        cursor = conn.cursor()
        query = """
            SELECT etc_file_url, date
            FROM submissions
            WHERE entity_name = %s
              AND date >= %s
              AND date <= %s
              AND etc_file_url IS NOT NULL
              AND TRIM(etc_file_url) <> ''
            ORDER BY date
        """
        cursor.execute(query, (plaza_name, from_date, to_date))
        rows = cursor.fetchall()

        if not rows:
            print(
                f"No ETC files found for entity_name='{plaza_name}' "
                f"between {from_date} and {to_date}."
            )
            return downloaded_paths

        print(f"Found {len(rows)} ETC file(s) to download.")

        for index, (file_url, file_date) in enumerate(rows, start=1):
            file_url = str(file_url).strip()
            if not file_url:
                continue

            filename = download_filename_from_url(file_url, file_date, index)
            target_path = output_folder / filename

            if target_path.exists():
                print(f"  [{index}/{len(rows)}] Already exists: {target_path.name}")
                downloaded_paths.append(target_path)
                continue

            print(f"  [{index}/{len(rows)}] Downloading: {file_url}")
            response = requests.get(file_url, timeout=120)
            response.raise_for_status()
            target_path.write_bytes(response.content)
            print(f"    Saved: {target_path}")
            downloaded_paths.append(target_path)

    finally:
        conn.close()

    return downloaded_paths


def write_tabular_file(file_path: Path, df: pd.DataFrame) -> None:
    """Write a DataFrame back to Excel or CSV based on file extension."""
    suffix = file_path.suffix.lower()
    if suffix == ".csv":
        df.to_csv(file_path, index=False)
        return
    if suffix in (".xlsx", ".xlsm", ".xls"):
        df.to_excel(file_path, index=False, engine="openpyxl")
        return
    raise ValueError(f"Unsupported file type for writing: {file_path}")


def load_etc_mop_lookup(
    etc_folder: Path,
    txn_id_column_names: list[str] | None = None,
) -> dict[str, str]:
    """
    Build transaction-id -> MOP lookup from all ETC files in the folder.
    IDs are normalized (letters/digits only; whitespace and special chars removed).
    """
    if txn_id_column_names is None:
        txn_id_column_names = get_etc_lookup_id_column_names()

    lookup: dict[str, str] = {}
    etc_files = sorted(
        p
        for p in etc_folder.iterdir()
        if p.is_file() and p.suffix.lower() in (".xlsx", ".xls", ".xlsm", ".csv")
    )

    if not etc_files:
        print(f"No ETC files found in {etc_folder}")
        return lookup

    print(f"\nBuilding lookup from {len(etc_files)} ETC file(s)...")

    for file_path in etc_files:
        df, header_row_idx = read_etc_tabular_file(file_path, txn_id_column_names)
        if df is None or df.empty:
            continue

        txn_id_col = find_column(df, txn_id_column_names)
        mop_col = find_column(df, MOP_COLUMN_NAMES)

        if not txn_id_col:
            print(
                f"  Skipping {file_path.name}: transaction ID column not found "
                f"for source '{ETC_LOOKUP_ID_COLUMN_SOURCE}' "
                f"within first {HEADER_SCAN_ROW_COUNT} row(s)."
            )
            continue
        if not mop_col:
            print(
                f"  Skipping {file_path.name}: MOP/Trip Type column not found "
                f"within first {HEADER_SCAN_ROW_COUNT} row(s)."
            )
            continue

        added = 0
        for _, row in df.iterrows():
            txn_id = normalize_txn_id_for_comparison(row[txn_id_col])
            if not txn_id or txn_id in lookup:
                continue

            mop_value = row[mop_col]
            if pd.isna(mop_value) or str(mop_value).strip() == "":
                lookup[txn_id] = NOT_FOUND_VALUE
            else:
                lookup[txn_id] = str(mop_value).strip()

            added += 1

        header_note = (
            f"header_row={header_row_idx + 1}"
            if header_row_idx is not None
            else "header_row=1 (default)"
        )
        print(
            f"  {file_path.name}: {header_note}, TxnId='{txn_id_col}', "
            f"MOP='{mop_col}', new lookup entries={added}"
        )

    print(f"Total unique transaction ID entries in lookup: {len(lookup)}")
    return lookup


def resolve_journey_type(cch_tran_id, lookup: dict[str, str]) -> str:
    """Return MOP for a normalized CCH Tran Id match, else NA."""
    normalized_id = normalize_txn_id_for_comparison(cch_tran_id)
    if not normalized_id:
        return NOT_FOUND_VALUE
    return lookup.get(normalized_id, NOT_FOUND_VALUE)


def process_report_file(file_path: Path, lookup: dict[str, str]) -> bool:
    """Add/update Journey Type in one report file using the ETC lookup."""
    df = read_tabular_file(file_path)
    if df is None:
        print(f"  Error reading {file_path.name}")
        return False

    cch_col = find_column(df, CCH_TRAN_ID_COLUMN_NAMES)
    if not cch_col:
        print(f"  Skipping {file_path.name}: 'CCH Tran Id' column not found.")
        return False

    if JOURNEY_TYPE_COLUMN not in df.columns:
        df[JOURNEY_TYPE_COLUMN] = None

    matched = 0
    not_found = 0
    empty_ids = 0

    for idx, row in df.iterrows():
        cch_value = row[cch_col]
        journey_type = resolve_journey_type(cch_value, lookup)
        df.at[idx, JOURNEY_TYPE_COLUMN] = journey_type

        if not normalize_txn_id_for_comparison(cch_value):
            empty_ids += 1
        elif journey_type == NOT_FOUND_VALUE:
            not_found += 1
        else:
            matched += 1

    try:
        write_tabular_file(file_path, df)
    except Exception as exc:
        print(f"  Error writing {file_path.name}: {exc}")
        return False

    print(
        f"  Updated {file_path.name}: column '{cch_col}' -> '{JOURNEY_TYPE_COLUMN}' "
        f"(matched={matched}, not_found={not_found}, empty_id={empty_ids})"
    )
    return True


def process_input_folder(input_folder: Path, lookup: dict[str, str]) -> tuple[int, int]:
    """Process all supported report files in the input folder."""
    if not input_folder.exists():
        raise FileNotFoundError(f"Input folder not found: {input_folder}")

    report_files = (
        sorted(input_folder.glob("*.xlsx"))
        + sorted(input_folder.glob("*.xls"))
        + sorted(input_folder.glob("*.xlsm"))
        + sorted(input_folder.glob("*.csv"))
    )
    if not report_files:
        print(f"No supported report files found in {input_folder}")
        return 0, 0

    print(f"\nProcessing {len(report_files)} report file(s) in {input_folder}...")
    success = 0
    failed = 0

    for file_path in report_files:
        if process_report_file(file_path, lookup):
            success += 1
        else:
            failed += 1

    return success, failed


def main():
    parse_date_yyyy_mm_dd(FROM_DATE)
    parse_date_yyyy_mm_dd(TO_DATE)

    print("=" * 80)
    print("ADD JOURNEY TYPE IN REPORT")
    print("=" * 80)
    print(f"Plaza: {PLAZA_NAME}")
    print(f"Date range: {FROM_DATE} to {TO_DATE}")
    print(f"Input folder: {INPUT_FOLDER}")
    print(f"ETC lookup ID column source: {ETC_LOOKUP_ID_COLUMN_SOURCE}")
    print(f"ETC header scan rows: {HEADER_SCAN_ROW_COUNT}")
    print("=" * 80)

    etc_folder = EXISTING_ETC_FOLDER.resolve() if EXISTING_ETC_FOLDER else get_etc_files_folder(PLAZA_NAME)
    print(f"\nETC files folder: {etc_folder}")

    if EXISTING_ETC_FOLDER:
        if not etc_folder.exists() or not etc_folder.is_dir():
            raise FileNotFoundError(f"ETC folder not found or not a directory: {etc_folder}")
        print("Skipping ETC download because EXISTING_ETC_FOLDER is set.")
    else:
        downloaded = download_etc_files(PLAZA_NAME, FROM_DATE, TO_DATE, etc_folder)
        print(f"Downloaded/available ETC files: {len(downloaded)}")

    lookup = load_etc_mop_lookup(etc_folder)
    if not lookup:
        print("\nWarning: ETC lookup is empty. All Journey Type values will be NA.")

    success, failed = process_input_folder(Path(INPUT_FOLDER), lookup)

    print("\n" + "=" * 80)
    print("SUMMARY")
    print("=" * 80)
    print(f"ETC files folder: {etc_folder}")
    print(f"Report files updated: {success}")
    print(f"Report files failed/skipped: {failed}")
    print("=" * 80)


if __name__ == "__main__":
    main()
