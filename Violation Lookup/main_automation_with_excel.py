from violation_extractor import extract_and_combine_excel_files
from update_exceptions import (
    clear_exceptions,
    remove_duplicates,
    update_exceptions_in_db,
)
from Violation_lookup import (
    process_single_excel_file,
    process_folder,
    match_and_generate_output,
)
from exception_value_count import clear_exception_values, calculate_exception_values
from update_manual_violation_from_excel import (
    clear_manual_violation,
    process_excel_and_update_db,
)
from pathlib import Path
from openpyxl import load_workbook
import pandas as pd

from db_config import get_process_date_range_strings

PREFER_CCH_TXN_ID_IN_LIFE_CYCLE_FILES = False


def build_combined_report_filename(
    plaza_name: str, from_date: str, to_date: str
    ) -> str:
    safe_plaza = plaza_name.replace(" ", "_").replace("/", "_")
    return f"{safe_plaza}_combined_{from_date}_to_{to_date}.xlsx"


def get_combined_report_path(
    plaza_name: str, from_date: str, to_date: str, base_folder: str = "."
    ) -> str:
    return str(
        Path(base_folder)
        / build_combined_report_filename(plaza_name, from_date, to_date)
    )


def mark_done_in_reference_excel(excel_file_path: str, row_index: int):
    # row_index is DataFrame index (0-based); +2 maps to Excel row (header on row 1)
    excel_row = row_index + 2

    wb = load_workbook(excel_file_path)
    sheet = wb.active

    # Find 'Done' column from header row
    done_col = None
    for col_idx in range(1, sheet.max_column + 1):
        if sheet.cell(row=1, column=col_idx).value == "Done":
            done_col = col_idx
            break

    if done_col is None:
        wb.close()
        raise ValueError("Column 'Done' not found in reference Excel file.")

    sheet.cell(row=excel_row, column=done_col, value="done")
    wb.save(excel_file_path)
    wb.close()
    print(
        f"Marked 'done' for Excel row {excel_row} in reference file: {excel_file_path}"
    )


def extract_violation_report(
    plaza_name: str, from_date: str, to_date: str, output_file: str = None
   ):
    print(f"\n{'=' * 80}")
    print("VIOLATION REPORT EXTRACTION")
    print(f"{'=' * 80}")
    print(f"Plaza Name: {plaza_name}")
    print(f"Date Range: {from_date} to {to_date}")

    # Generate output filename if not provided
    if not output_file:
        output_file = build_combined_report_filename(plaza_name, from_date, to_date)
        print(f"Output File: {output_file}")

    print(f"{'=' * 80}\n")

    # Extract and combine Excel files using date range (not specific_dates)
    output_path = extract_and_combine_excel_files(
        entity_name=plaza_name,
        from_date=from_date,
        to_date=to_date,
        output_file=output_file,
        specific_dates=None,  # Exclude specific dates logic, use date range only
    )

    if output_path:
        print(f"\n{'=' * 80}")
        print("EXTRACTION COMPLETE")
        print(f"{'=' * 80}")
        print(f"Output file: {output_path}")
        print(f"{'=' * 80}")
        return output_path
    else:
        print(f"\n{'=' * 80}")
        print("EXTRACTION FAILED")
        print(f"{'=' * 80}")
        return None


def process_violation_lookup(
    excel_file_path: str,
    folder_path: str,
    prefer_cch_txn_id_in_life_cycle_files: bool = False,
     ):
    print(f"\n{'=' * 80}")
    print("VIOLATION LOOKUP PROCESSING")
    print(f"{'=' * 80}")
    print(f"Excel File: {excel_file_path}")
    print(f"Folder: {folder_path}")
    print(f"{'=' * 80}")

    # Check if Excel file exists
    if not Path(excel_file_path).exists():
        print(f"\nError: Excel file not found: {excel_file_path}")
        return False

    # Check if folder exists
    if not Path(folder_path).exists():
        print(f"\nError: Folder not found: {folder_path}")
        return False

    # Process single Excel file
    print(f"\n{'=' * 80}")
    print("STEP 1: Processing single Excel file")
    print(f"{'=' * 80}")
    single_file_results, df_single = process_single_excel_file(excel_file_path)

    if not single_file_results or df_single is None:
        print("Error: Could not process single Excel file.")
        return False

    # Process all Excel files in the folder
    print(f"\n{'=' * 80}")
    print("STEP 2: Processing Excel files in folder")
    print(f"{'=' * 80}")
    if prefer_cch_txn_id_in_life_cycle_files:
        print("Lifecycle lookup priority: CCH Txn Id first, then Agency Txn Id")
    else:
        print("Lifecycle lookup priority: Agency Txn Id first, then CCH Txn Id")
    folder_results, violation_amts, reader_read_times, journey_types = process_folder(
        folder_path,
        prefer_cch_txn_id=prefer_cch_txn_id_in_life_cycle_files,
    )

    if not folder_results:
        print("Warning: No results from folder files. Skipping comparison.")
        return False

    # Match and mark violation_matched on the combined report (no intermediate output files)
    print(f"\n{'=' * 80}")
    print("STEP 3: Matching & updating combined report")
    print(f"{'=' * 80}")
    print(
        "Flow: Compare single file rows (CCH TXN NO) with folder data (Agency Txn Id)"
    )
    print(
        "  -> IDs normalized before compare (whitespace/backticks/special chars removed)"
    )
    print(
        "  -> For ALL matches: Write 'Journey Type', 'Reader Read Time', and 'Violation Amt' to combined Excel"
    )
    print(
        "  -> If Violation Amt is non-zero -> mark 'Yes' in violation_matched"
    )
    print("  -> Otherwise -> mark 'No' in violation_matched")
    print(
        "  -> Does not generate violation_lookup_output files (DB update uses combined report)"
    )

    success = match_and_generate_output(
        excel_file_path,
        df_single,
        folder_path,
        folder_results,
        violation_amts,
        reader_read_times,
        journey_types,
        write_matched_output=False,
    )

    if success:
        print(f"\n{'=' * 80}")
        print("VIOLATION LOOKUP COMPLETE")
        print(f"{'=' * 80}")
        print(f"Updated Excel file: {excel_file_path}")
        print(f"{'=' * 80}")
        return True
    else:
        print(f"\n{'=' * 80}")
        print("VIOLATION LOOKUP FAILED")
        print(f"{'=' * 80}")
        return False


def process_plazas_from_excel(
    excel_file_path: str,
    life_cycle_base_folder: str,
    resume_from_violation_lookup: bool = False,
    combined_report_base_folder: str = ".",
   ):
    try:
        from_date, to_date = get_process_date_range_strings()
    except (RuntimeError, ValueError) as exc:
        print(f"\nError: {exc}")
        return

    print(f"\nUsing process date range from .env: {from_date} to {to_date}")

    # Load Excel file into a DataFrame
    df = pd.read_excel(excel_file_path)

    if resume_from_violation_lookup:
        print(f"\n{'=' * 80}")
        print("RESUME MODE ENABLED")
        print(f"{'=' * 80}")
        print("Skipping extraction, duplicate removal, and exception count update.")
        print(
            "Program will resume from violation lookup using the combined report name"
        )
        print("derived from plaza name and PROCESS_FROM_DATE / PROCESS_TO_DATE in .env.")
        print(f"{'=' * 80}")

    # Iterate over each row in the DataFrame
    for idx, row in df.iterrows():
        # Check if the 'Done' column is empty
        if pd.isna(row["Done"]):
            plaza_name_db = (
                str(row["Plaza_name_DB"]).strip()
                if pd.notna(row["Plaza_name_DB"])
                else ""
            )
            plaza_name_idfc = (
                str(row["Plaza_name_IDFC"]).strip()
                if pd.notna(row["Plaza_name_IDFC"])
                else ""
            )

            # Construct folder_path from configurable base folder
            folder_path = str(Path(life_cycle_base_folder) / plaza_name_idfc)

            # Process the plaza
            print(f"\nProcessing Plaza: {plaza_name_db}")

            if resume_from_violation_lookup:
                output_path = get_combined_report_path(
                    plaza_name_db,
                    from_date,
                    to_date,
                    combined_report_base_folder,
                )
                print(f"Resume file path assumed: {output_path}")
                if not Path(output_path).exists():
                    print(f"Combined report file not found for resume: {output_path}")
                    continue
            else:
                output_path = extract_violation_report(
                    plaza_name_db, from_date, to_date
                )

                if output_path:
                    print(f"Extraction complete for {plaza_name_db}")
                    print(f"Extraction report path: {output_path}")
                    # Step 2: Clear old exceptions values
                    print(f"Clearing exception counts for {plaza_name_db} in DB")
                    clear_exceptions(plaza_name_db, from_date, to_date)
                    print(f"Exception counts are cleared for {plaza_name_db} in DB")

                    # Step 3: Remove duplicates based on 'CCH TXN NO'
                    print(
                        f"Removing duplicates for {plaza_name_db} based on 'CCH TXN NO'"
                    )
                    plaza_df = remove_duplicates(output_path)
                    print(
                        f"Duplicates removed for {plaza_name_db} based on 'CCH TXN NO'"
                    )

                    # Step 4: Update exceptions in DB for each (date, shift) combination
                    print(f"Updating exception values for {plaza_name_db} in DB")
                    update_exceptions_in_db(plaza_df, plaza_name_db)
                    print(f"Exceptions updated for {plaza_name_db} in DB")
                else:
                    print(f"Report extraction failed for {plaza_name_db}.")
                    continue

            # Step 5: Process violation lookup
            print(
                f"Processing violation lookup with Life Cycle files for {plaza_name_db}"
            )
            violation_lookup_success = process_violation_lookup(
                output_path,
                folder_path,
                prefer_cch_txn_id_in_life_cycle_files=PREFER_CCH_TXN_ID_IN_LIFE_CYCLE_FILES,
            )

            if violation_lookup_success:
                print(f"Violation lookup complete for {plaza_name_db}")

                print(f"Clearing exception values for {plaza_name_db} in DB")
                clear_exception_values(plaza_name_db, from_date, to_date)
                print(f"Exception values cleared for {plaza_name_db}")

                print(f"Calculating exception values for {plaza_name_db}")
                calculate_exception_values(output_path, plaza_name_db)
                print(f"Exceptions calculated and updated in DB for {plaza_name_db}")

                # Clear old manual violation values first
                print(f"Clearing manual violation values for {plaza_name_db}")
                clear_result = clear_manual_violation(plaza_name_db, from_date, to_date)
                print(f"Manual violation values cleared for {plaza_name_db}")

                if clear_result.get("errors", 0) > 0:
                    print("\n✗ Error during manual violation cleanup. Skipping update.")
                    print(
                        f"Error: {clear_result.get('errors', 0)} errors occurred during manual violation cleanup."
                    )
                    print(f"Plaza {plaza_name_db} processing failed.")
                else:
                    print("\n✓ Cleanup complete. Proceeding to update...\n")

                    # Process the same input file (output_path) that was generated in step 1
                    print(
                        f"Processing manual violation values from input file for {plaza_name_db}"
                    )
                    stats = process_excel_and_update_db(output_path, plaza_name_db)
                    print(f"Manual violation values updated for {plaza_name_db}")
                    if stats:
                        print(
                            f"  Updated: {stats.get('updated', 0)}, Skipped: {stats.get('skipped', 0)}"
                        )
                        # Mark row done only after manual violation update succeeds
                        df.at[idx, "Done"] = "done"
                        mark_done_in_reference_excel(excel_file_path, idx)
                        print(
                            f"{plaza_name_db} processed successfully. Marking as 'done' in Excel file."
                        )

            else:
                print(f"Violation lookup failed for {plaza_name_db}.")


if __name__ == "__main__":
    excel_file_path = "Reference_File.xlsx"
    # life_cycle_base_folder = "ICICI_files"
    life_cycle_base_folder = "IDFC_Life_cycle_files_from_July07_to_August09"
    resume_from_violation_lookup = False
    combined_report_base_folder = "."

    process_plazas_from_excel(
        excel_file_path,
        life_cycle_base_folder,
        resume_from_violation_lookup=resume_from_violation_lookup,
        combined_report_base_folder=combined_report_base_folder,
    )
