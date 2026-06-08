"""
YES Bank Portal Login Bot

Current scope:
1) Read reference Excel with headers: Plaza_name, Username, Password, Done
2) Open YES Bank portal login page
3) Enter username and password for each pending row
4) Click Submit
5) Click Home
6) Open Revenue Report from MIS Reports
7) Select the only available client from dropdown
8) Submit configured date range in month-sized intervals
9) Export Debit Adjustment reports where Debit Adjustment Count > 1
10) Rename downloads using plaza name + settlement date and move them into plaza folder
11) Mark the row as done in Excel after all intervals are submitted
"""

from __future__ import annotations

import shutil
import time
from datetime import datetime, timedelta
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from selenium import webdriver
from selenium.common.exceptions import StaleElementReferenceException, WebDriverException
from selenium.webdriver.common.action_chains import ActionChains
from selenium.webdriver.common.by import By
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.support.ui import Select
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.support.ui import WebDriverWait


LOGIN_URL = "https://netcybl.maximusinfoware.in/frmLogin.aspx"
REFERENCE_EXCEL = "YES_Reference_file.xlsx"
DOWNLOADS_FOLDER = r"C:\Users\Divyesh Parmar\Downloads"
MAX_ELEMENT_WAIT_SECONDS = 120
MAX_DOWNLOAD_WAIT_SECONDS = 120
POST_LOGIN_WAIT_SECONDS = 8
MAX_INTERVAL_ATTEMPTS = 5
MAX_CALENDAR_SELECT_ATTEMPTS = 5
POLL_SECONDS = 2
DATE_RANGE_START = "2025-11-01"  # YYYY-MM-DD
DATE_RANGE_END = "2026-05-31"    # YYYY-MM-DD
SCRIPT_DIR = Path(__file__).resolve().parent


def setup_driver(headless: bool = False):
    """Create a Chrome driver with practical defaults."""
    chrome_options = Options()
    if headless:
        chrome_options.add_argument("--headless")
    else:
        chrome_options.add_argument("--start-maximized")
    chrome_options.add_argument("--no-sandbox")
    chrome_options.add_argument("--disable-dev-shm-usage")
    chrome_options.add_argument("--disable-blink-features=AutomationControlled")
    chrome_options.add_experimental_option("excludeSwitches", ["enable-automation"])
    chrome_options.add_experimental_option("useAutomationExtension", False)
    prefs = {
        "download.default_directory": str(Path(DOWNLOADS_FOLDER).resolve()),
        "download.prompt_for_download": False,
        "download.directory_upgrade": True,
        "safebrowsing.enabled": True,
    }
    chrome_options.add_experimental_option("prefs", prefs)
    driver = webdriver.Chrome(options=chrome_options)
    if not headless:
        driver.maximize_window()
    return driver


def normalize_header_name(value) -> str:
    """Normalize Excel headers for case-insensitive matching."""
    if value is None:
        return ""
    return str(value).strip().lower()


def resolve_required_columns(df: pd.DataFrame) -> dict:
    """Resolve Plaza_name, Username, Password, Done columns case-insensitively."""
    wanted = {
        "plaza_name": None,
        "username": None,
        "password": None,
        "done": None,
    }

    for col in df.columns:
        norm = normalize_header_name(col)
        if norm in wanted and wanted[norm] is None:
            wanted[norm] = col

    missing = [name for name, col in wanted.items() if col is None]
    if missing:
        raise ValueError(
            f"Missing required columns in {REFERENCE_EXCEL}: {missing}. "
            "Expected headers: Plaza_name, Username, Password, Done (case-insensitive)."
        )

    return wanted


def mark_done_in_excel(excel_file_path: str, row_index: int, done_header_name):
    """Write Done='done' for the given DataFrame row index."""
    excel_row = row_index + 2
    workbook = load_workbook(excel_file_path)
    sheet = workbook.active

    done_col = None
    for col_idx in range(1, sheet.max_column + 1):
        header = sheet.cell(row=1, column=col_idx).value
        if normalize_header_name(header) == normalize_header_name(done_header_name):
            done_col = col_idx
            break

    if done_col is None:
        workbook.close()
        raise ValueError("Could not locate 'Done' column in Excel while marking status.")

    sheet.cell(row=excel_row, column=done_col, value="done")
    workbook.save(excel_file_path)
    workbook.close()
    print(f"Marked row {excel_row} as done in {excel_file_path}")


def wait_for_visible_element(driver, by, value, timeout_seconds: int = MAX_ELEMENT_WAIT_SECONDS):
    """Wait until an element is visible and return it."""
    return WebDriverWait(driver, timeout_seconds).until(
        EC.visibility_of_element_located((by, value))
    )


def click_element_by_id(driver, element_id: str):
    """Click an element by id via JavaScript (avoids stale WebElement references)."""
    clicked = driver.execute_script(
        """
        const element = document.getElementById(arguments[0]);
        if (!element) {
            return false;
        }
        element.click();
        return true;
        """,
        element_id,
    )
    if not clicked:
        raise RuntimeError(f"Could not click element with id={element_id!r}.")


def wait_for_update_panel_idle(driver, timeout_seconds: int = MAX_ELEMENT_WAIT_SECONDS):
    """Wait until an Ajax UpdatePanel refresh finishes (spinner hidden, date inputs back)."""
    def panel_idle(web_driver) -> bool:
        try:
            return bool(
                web_driver.execute_script(
                    """
                    const progress = document.getElementById('MainContentPlaceHolder_UpdateProgress1');
                    if (progress) {
                        const style = (progress.getAttribute('style') || '').toLowerCase();
                        const hidden =
                            style.includes('display: none') || style.includes('display:none');
                        if (!hidden && progress.offsetParent !== null) {
                            return false;
                        }
                    }

                    const fromInput = document.getElementById(arguments[0]);
                    const toInput = document.getElementById(arguments[1]);
                    return !!fromInput && fromInput.offsetParent !== null
                        && !!toInput && toInput.offsetParent !== null;
                    """,
                    FROM_DATE_INPUT_ID,
                    TO_DATE_INPUT_ID,
                )
            )
        except Exception:
            return False

    WebDriverWait(driver, timeout_seconds).until(panel_idle)


def trigger_date_postback(driver, event_target: str):
    """Run the portal's partial postback for a date field (syncs ASP.NET ViewState)."""
    triggered = driver.execute_script(
        """
        if (typeof __doPostBack !== 'function') {
            return false;
        }
        __doPostBack(arguments[0], '');
        return true;
        """,
        event_target,
    )
    if not triggered:
        raise RuntimeError(f"Could not trigger date postback for {event_target!r}.")
    wait_for_update_panel_idle(driver)


def wait_for_show_postback(driver, timeout_seconds: int = MAX_ELEMENT_WAIT_SECONDS):
    """
    Wait for the Ajax UpdatePanel refresh that runs after Show is clicked.
    Uses document queries only so partial DOM updates do not stale WebElements.
    """
    def postback_complete(web_driver) -> bool:
        try:
            return bool(
                web_driver.execute_script(
                    """
                    const progress = document.getElementById('MainContentPlaceHolder_UpdateProgress1');
                    if (progress) {
                        const style = (progress.getAttribute('style') || '').toLowerCase();
                        const hidden =
                            style.includes('display: none') || style.includes('display:none');
                        if (!hidden && progress.offsetParent !== null) {
                            return false;
                        }
                    }

                    const fromInput = document.getElementById(arguments[0]);
                    if (!fromInput || fromInput.offsetParent === null) {
                        return false;
                    }

                    const grid = document.getElementById('MainContentPlaceHolder_grdView');
                    if (grid && grid.offsetParent !== null) {
                        return true;
                    }

                    const bodyText = (document.body && document.body.innerText || '').toLowerCase();
                    return bodyText.includes('no record');
                    """,
                    FROM_DATE_INPUT_ID,
                )
            )
        except Exception:
            return False

    WebDriverWait(driver, timeout_seconds).until(postback_complete)


def is_502_gateway_page(driver) -> bool:
    """Return True when the browser shows the known 502 proxy error page."""
    checks = [
        "502 - Web server received an invalid response while acting as a gateway or proxy server.",
        "received an invalid response from the content server",
    ]

    try:
        page_text = (driver.find_element(By.TAG_NAME, "body").text or "").strip().lower()
    except Exception:
        page_text = ""

    if page_text and all(check.lower() in page_text for check in checks):
        return True

    try:
        fieldsets = driver.find_elements(By.TAG_NAME, "fieldset")
    except Exception:
        return False

    for fieldset in fieldsets:
        try:
            text = (fieldset.text or "").strip().lower()
        except Exception:
            continue
        if text and all(check.lower() in text for check in checks):
            return True

    return False


def ensure_not_502_page(driver):
    """Raise when the current page is the known 502 proxy error page."""
    if is_502_gateway_page(driver):
        raise RuntimeError("Detected 502 gateway/proxy error page.")


def collect_visible_messages(driver) -> list[str]:
    """Best-effort collection of visible error or status messages."""
    selectors = [
        ".alert",
        ".error",
        ".text-danger",
        ".validation-summary-errors",
        "[class*='error']",
        "[id*='error']",
        "[id*='lbl']",
    ]

    messages: list[str] = []
    seen: set[str] = set()

    for selector in selectors:
        try:
            elements = driver.find_elements(By.CSS_SELECTOR, selector)
        except Exception:
            continue

        for elem in elements:
            try:
                if not elem.is_displayed():
                    continue
                text = (elem.text or "").strip()
                if text and text not in seen:
                    seen.add(text)
                    messages.append(text)
            except Exception:
                continue

    return messages


def is_login_form_visible(driver) -> bool:
    """Return True when the username/password fields are still visible."""
    try:
        username_fields = driver.find_elements(By.ID, "txtUserName")
        password_fields = driver.find_elements(By.ID, "txtPassword")
        for field in username_fields + password_fields:
            if field.is_displayed():
                return True
    except Exception:
        pass
    return False


def login_once(driver, username: str, password: str) -> bool:
    """Open login page, submit credentials, and determine whether login succeeded."""
    driver.get(LOGIN_URL)
    ensure_not_502_page(driver)

    username_field = wait_for_visible_element(driver, By.ID, "txtUserName")
    password_field = wait_for_visible_element(driver, By.ID, "txtPassword")
    submit_button = wait_for_visible_element(driver, By.ID, "btnLogin")

    username_field.clear()
    username_field.send_keys(username)
    password_field.clear()
    password_field.send_keys(password)
    submit_button.click()

    deadline = time.time() + POST_LOGIN_WAIT_SECONDS
    while time.time() < deadline:
        ensure_not_502_page(driver)
        if not is_login_form_visible(driver):
            return True
        time.sleep(1)

    return False


def parse_date_yyyy_mm_dd(value: str) -> datetime:
    """Parse program date variable in YYYY-MM-DD format."""
    return datetime.strptime(value.strip(), "%Y-%m-%d")


def parse_portal_settlement_date(value: str) -> datetime:
    """Parse settlement date text shown in the portal table."""
    return datetime.strptime(value.strip(), "%d/%b/%Y")


def format_dd_mm_yyyy(value: datetime) -> str:
    """Format date for portal inputs."""
    return value.strftime("%d/%m/%Y")


def parse_dd_mm_yyyy(value: str) -> datetime:
    """Parse portal date input text in DD/MM/YYYY format."""
    return datetime.strptime(value.strip(), "%d/%m/%Y")


def format_calendar_day_title(value: datetime) -> str:
    """Format a date the way Ajax CalendarBehavior titles day cells."""
    return value.strftime("%A, %B %d, %Y")


FROM_DATE_INPUT_ID = "MainContentPlaceHolder_txtFromDate"
TO_DATE_INPUT_ID = "MainContentPlaceHolder_txtToDate"
FROM_CALENDAR_BEHAVIOR_ID = "MainContentPlaceHolder_cetxtFromDate"
TO_CALENDAR_BEHAVIOR_ID = "MainContentPlaceHolder_cetxtToDate"
FROM_DATE_POSTBACK_TARGET = "ctl00$MainContentPlaceHolder$txtFromDate"
TO_DATE_POSTBACK_TARGET = "ctl00$MainContentPlaceHolder$txtToDate"


def sanitize_filename_part(value: str) -> str:
    """Sanitize text for safe filenames and folder names."""
    safe = "".join(ch if ch.isalnum() or ch in ("-", "_") else "_" for ch in str(value).strip())
    return safe.strip("_") or "unknown"


def get_plaza_output_dir(plaza_name: str) -> Path:
    """Create and return a plaza-named folder next to the script."""
    plaza_dir = SCRIPT_DIR / sanitize_filename_part(plaza_name)
    plaza_dir.mkdir(parents=True, exist_ok=True)
    return plaza_dir


def snapshot_download_folder(downloads_folder: str = DOWNLOADS_FOLDER) -> dict[str, float]:
    """Capture existing download files and their mtimes."""
    folder = Path(downloads_folder)
    snapshot: dict[str, float] = {}
    if not folder.exists():
        return snapshot

    for pattern in ("*.xls", "*.xlsx", "*.csv"):
        for path in folder.glob(pattern):
            try:
                snapshot[str(path)] = path.stat().st_mtime
            except Exception:
                continue

    return snapshot


def wait_for_new_download(
    previous_snapshot: dict[str, float],
    downloads_folder: str = DOWNLOADS_FOLDER,
    timeout_seconds: int = MAX_DOWNLOAD_WAIT_SECONDS,
) -> Path:
    """Wait for a newly completed download file to appear in Downloads."""
    folder = Path(downloads_folder)
    if not folder.exists():
        raise RuntimeError(f"Downloads folder not found: {downloads_folder}")

    deadline = time.time() + timeout_seconds
    latest_candidate = None

    while time.time() < deadline:
        active_downloads = list(folder.glob("*.crdownload"))
        candidates: list[Path] = []

        for pattern in ("*.xls", "*.xlsx", "*.csv"):
            candidates.extend(folder.glob(pattern))

        fresh_candidates: list[Path] = []
        for path in candidates:
            try:
                mtime = path.stat().st_mtime
            except Exception:
                continue

            old_mtime = previous_snapshot.get(str(path))
            if old_mtime is None or mtime > old_mtime:
                fresh_candidates.append(path)

        if fresh_candidates and not active_downloads:
            fresh_candidates.sort(key=lambda item: item.stat().st_mtime, reverse=True)
            latest_candidate = fresh_candidates[0]
            break

        time.sleep(POLL_SECONDS)

    if latest_candidate is None:
        raise RuntimeError(f"Downloaded file not found within {timeout_seconds} seconds.")

    return latest_candidate


def move_download_to_plaza_folder(downloaded_file: Path, plaza_name: str, settlement_date_text: str) -> Path:
    """Rename the downloaded file and move it into the plaza folder."""
    plaza_dir = get_plaza_output_dir(plaza_name)
    settlement_date = parse_portal_settlement_date(settlement_date_text).strftime("%Y-%m-%d")
    safe_plaza = sanitize_filename_part(plaza_name)
    target_path = plaza_dir / f"{safe_plaza}_{settlement_date}{downloaded_file.suffix.lower()}"

    if target_path.exists():
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        target_path = plaza_dir / f"{safe_plaza}_{settlement_date}_{stamp}{downloaded_file.suffix.lower()}"

    shutil.move(str(downloaded_file), str(target_path))
    print(f"Moved download: {target_path}")
    return target_path


def split_date_range_to_month_intervals(start_date: str, end_date: str) -> list[tuple[datetime, datetime]]:
    """
    Split an inclusive date range into month-bounded intervals.
    A full 31-day month remains a 31-day interval.
    """
    start_dt = parse_date_yyyy_mm_dd(start_date)
    end_dt = parse_date_yyyy_mm_dd(end_date)
    if end_dt < start_dt:
        raise ValueError("DATE_RANGE_END must be >= DATE_RANGE_START")

    intervals: list[tuple[datetime, datetime]] = []
    current = start_dt

    while current <= end_dt:
        if current.month == 12:
            next_month_start = datetime(current.year + 1, 1, 1)
        else:
            next_month_start = datetime(current.year, current.month + 1, 1)
        month_end = next_month_start - timedelta(days=1)
        interval_end = min(month_end, end_dt)
        intervals.append((current, interval_end))
        current = interval_end + timedelta(days=1)

    return intervals


def navigate_to_revenue_report(driver):
    """After login, click Home and open Revenue Report from MIS Reports."""
    ensure_not_502_page(driver)
    home_link = wait_for_visible_element(
        driver,
        By.XPATH,
        "//a[@title='Home Page' and contains(normalize-space(.), 'Home')]",
    )
    home_link.click()

    mis_reports_link = wait_for_visible_element(
        driver,
        By.XPATH,
        "//a[normalize-space()='MIS Reports']",
    )
    ActionChains(driver).move_to_element(mis_reports_link).perform()
    time.sleep(1)

    revenue_report_link = wait_for_visible_element(
        driver,
        By.XPATH,
        "//a[contains(@href, 'RevenueReport.aspx') and normalize-space()='Revenue Report']",
    )
    try:
        revenue_report_link.click()
    except Exception:
        driver.execute_script("arguments[0].click();", revenue_report_link)

    WebDriverWait(driver, MAX_ELEMENT_WAIT_SECONDS).until(
        lambda d: "RevenueReport.aspx" in d.current_url
    )
    ensure_not_502_page(driver)


def select_only_client_option(driver):
    """Select the only non-placeholder client option from the dropdown."""
    ensure_not_502_page(driver)
    dropdown_element = wait_for_visible_element(driver, By.ID, "MainContentPlaceHolder_ddlClient")
    select = Select(dropdown_element)
    valid_options = [option for option in select.options if (option.get_attribute("value") or "").strip() != "0"]

    if len(valid_options) != 1:
        raise ValueError(f"Expected exactly one selectable client option, found {len(valid_options)}.")

    selected_value = valid_options[0].get_attribute("value")
    select.select_by_value(selected_value)
    time.sleep(2)


def disable_date_input_postbacks(driver):
    """
    Prevent From/To onchange postbacks while automation sets dates.
    Changing From Date alone triggers a partial postback that resets To Date.
    """
    driver.execute_script(
        """
        [arguments[0], arguments[1]].forEach(function (inputId) {
            const input = document.getElementById(inputId);
            if (!input) {
                return;
            }
            input.onchange = null;
            input.onblur = null;
            input.onkeyup = null;
            input.removeAttribute('onchange');
            input.removeAttribute('onblur');
            input.removeAttribute('onkeyup');
        });
        """,
        FROM_DATE_INPUT_ID,
        TO_DATE_INPUT_ID,
    )


def read_revenue_report_date_state(driver) -> dict[str, str]:
    """Read visible input values and Ajax calendar selected dates."""
    return driver.execute_script(
        """
        function formatDate(value) {
            if (!(value instanceof Date) || isNaN(value)) {
                return '';
            }
            const day = String(value.getDate()).padStart(2, '0');
            const month = String(value.getMonth() + 1).padStart(2, '0');
            const year = value.getFullYear();
            return `${day}/${month}/${year}`;
        }

        function readBehavior(behaviorId) {
            const behavior = window.$find ? window.$find(behaviorId) : null;
            if (!behavior) {
                return '';
            }
            const selected =
                typeof behavior.get_selectedDate === 'function'
                    ? behavior.get_selectedDate()
                    : behavior._selectedDate;
            return formatDate(selected);
        }

        const fromInput = document.getElementById(arguments[0]);
        const toInput = document.getElementById(arguments[1]);
        return {
            fromInput: fromInput ? (fromInput.value || '').trim() : '',
            toInput: toInput ? (toInput.value || '').trim() : '',
            fromBehavior: readBehavior(arguments[2]),
            toBehavior: readBehavior(arguments[3]),
        };
        """,
        FROM_DATE_INPUT_ID,
        TO_DATE_INPUT_ID,
        FROM_CALENDAR_BEHAVIOR_ID,
        TO_CALENDAR_BEHAVIOR_ID,
    )


def assert_date_state_matches(
    state: dict[str, str],
    expected_from: str,
    expected_to: str,
    stage: str,
    *,
    verify_to: bool = True,
) -> None:
    """Raise when input or calendar behavior dates do not match expectations."""
    from_input = state.get("fromInput", "")
    to_input = state.get("toInput", "")
    from_behavior = state.get("fromBehavior", "")
    to_behavior = state.get("toBehavior", "")

    mismatches: list[str] = []
    if from_input != expected_from:
        mismatches.append(f"From input expected {expected_from}, got {from_input}")
    if from_behavior and from_behavior != expected_from:
        mismatches.append(f"From calendar expected {expected_from}, got {from_behavior}")

    if verify_to:
        if to_input != expected_to:
            mismatches.append(f"To input expected {expected_to}, got {to_input}")
        if to_behavior and to_behavior != expected_to:
            mismatches.append(f"To calendar expected {expected_to}, got {to_behavior}")
        if expected_from != expected_to and to_input == from_input:
            mismatches.append(
                f"To date collapsed to From date ({to_input}) before {stage}"
            )

    if mismatches:
        raise RuntimeError(f"Date verification failed at {stage}: " + "; ".join(mismatches))


def _calendar_behavior_js() -> str:
    """Shared JavaScript helpers for Ajax CalendarBehavior."""
    return """
        function parseDdMmYyyy(value) {
            const parts = value.split('/');
            if (parts.length !== 3) {
                return null;
            }
            const day = parseInt(parts[0], 10);
            const month = parseInt(parts[1], 10) - 1;
            const year = parseInt(parts[2], 10);
            return new Date(year, month, day, 12, 0, 0, 0);
        }

        function applyBehaviorDate(behavior, selectedDate, value) {
            if (!behavior || !selectedDate) {
                return;
            }
            if (typeof behavior.set_selectedDate === 'function') {
                behavior.set_selectedDate(selectedDate);
            } else {
                behavior._selectedDate = selectedDate;
            }
            if (typeof behavior.set_visibleDate === 'function') {
                behavior.set_visibleDate(selectedDate);
            } else {
                behavior._visibleDate = selectedDate;
            }
            if (behavior._textbox && 'value' in behavior._textbox) {
                behavior._textbox.value = value;
                if ('defaultValue' in behavior._textbox) {
                    behavior._textbox.defaultValue = value;
                }
            }
        }
    """


def set_calendar_behavior_date(
    driver,
    input_id: str,
    behavior_id: str,
    target_dt: datetime,
) -> bool:
    """Set a date through CalendarBehavior APIs and the bound textbox."""
    target_value = format_dd_mm_yyyy(target_dt)
    return bool(
        driver.execute_script(
            _calendar_behavior_js()
            + """
            const input = document.getElementById(arguments[0]);
            const behavior = window.$find ? window.$find(arguments[1]) : null;
            const value = arguments[2];
            const selectedDate = parseDdMmYyyy(value);
            if (!input || !selectedDate) {
                return false;
            }

            applyBehaviorDate(behavior, selectedDate, value);
            input.value = value;
            input.defaultValue = value;
            if (typeof input.setAttribute === 'function') {
                input.setAttribute('value', value);
            }
            return (input.value || '').trim() === value;
            """,
            input_id,
            behavior_id,
            target_value,
        )
    )


def js_get_visible_calendar_title(driver, calendar_prefix: str) -> str:
    """Read the month title only when that calendar's day grid is visible."""
    return (
        driver.execute_script(
            """
            const prefix = arguments[0];
            const days = document.getElementById(prefix + '_days');
            if (!days || days.offsetParent === null) {
                return '';
            }
            const title = document.getElementById(prefix + '_title');
            return title ? (title.textContent || '').trim() : '';
            """,
            calendar_prefix,
        )
        or ""
    )


def set_calendar_visible_month(
    driver,
    behavior_id: str,
    target_dt: datetime,
    ) -> None:
    """Jump the open calendar popup to the month that contains target_dt."""
    month_anchor = datetime(target_dt.year, target_dt.month, 1)
    anchor_value = format_dd_mm_yyyy(month_anchor)
    driver.execute_script(
        _calendar_behavior_js()
        + """
        const behavior = window.$find ? window.$find(arguments[0]) : null;
        const selectedDate = parseDdMmYyyy(arguments[1]);
        if (!behavior || !selectedDate) {
            return false;
        }
        if (typeof behavior.set_visibleDate === 'function') {
            behavior.set_visibleDate(selectedDate);
        } else {
            behavior._visibleDate = selectedDate;
        }
        return true;
        """,
        behavior_id,
        anchor_value,
    )


def _parse_calendar_title_month_year(title_text: str) -> tuple[int, int]:
    """Parse calendar header text like 'February, 2026'."""
    cleaned = (title_text or "").strip().replace(",", "")
    parts = cleaned.split()
    if len(parts) < 2:
        raise ValueError(f"Unrecognized calendar title: {title_text!r}")

    month_name = parts[0]
    year = int(parts[-1])
    month = datetime.strptime(month_name, "%B").month
    return month, year


def js_get_element_text_by_id(driver, element_id: str) -> str:
    """Read element text via JavaScript (avoids stale WebElement references)."""
    return (
        driver.execute_script(
            """
            const element = document.getElementById(arguments[0]);
            return element ? (element.textContent || '').trim() : '';
            """,
            element_id,
        )
        or ""
    )


def js_get_input_value_by_id(driver, input_id: str) -> str:
    """Read an input value via JavaScript (avoids stale WebElement references)."""
    return (
        driver.execute_script(
            """
            const element = document.getElementById(arguments[0]);
            return element ? (element.value || '').trim() : '';
            """,
            input_id,
        )
        or ""
    )


def js_is_calendar_days_visible(driver, calendar_prefix: str) -> bool:
    """Return True when the Ajax calendar day grid is visible."""
    return bool(
        driver.execute_script(
            """
            const days = document.getElementById(arguments[0] + '_days');
            return !!(days && days.offsetParent !== null);
            """,
            calendar_prefix,
        )
    )


def is_browser_session_lost(exc: Exception) -> bool:
    """Return True when the Chrome/WebDriver session is no longer usable."""
    text = str(exc).lower()
    markers = (
        "invalid session id",
        "session deleted",
        "browser has closed",
        "not connected to devtools",
        "chrome not reachable",
        "disconnected: not connected",
    )
    return any(marker in text for marker in markers)


def dismiss_active_calendar_popup(driver):
    """
    Close a visible calendar popup using CalendarBehavior APIs.

    Do not force-hide .ajax__calendar nodes in the DOM; that breaks the embedded
    From calendar widget and prevents the To calendar from opening correctly.
    """
    driver.execute_script(
        """
        [arguments[0], arguments[1]].forEach(function (behaviorId) {
            const behavior = window.$find ? window.$find(behaviorId) : null;
            if (!behavior) {
                return;
            }
            if (typeof behavior.hide === 'function') {
                behavior.hide();
            } else if (typeof behavior.set_visible === 'function') {
                behavior.set_visible(false);
            }
        });

        if (document.activeElement && typeof document.activeElement.blur === 'function') {
            document.activeElement.blur();
        }
        """,
        FROM_CALENDAR_BEHAVIOR_ID,
        TO_CALENDAR_BEHAVIOR_ID,
    )


def navigate_calendar_to_month(
    driver,
    calendar_prefix: str,
    behavior_id: str,
    target_dt: datetime,
     ) -> bool:
    """Move the visible Ajax calendar popup to the month that contains target_dt."""
    prev_id = f"{calendar_prefix}_prevArrow"
    next_id = f"{calendar_prefix}_nextArrow"
    target_month = target_dt.month
    target_year = target_dt.year

    set_calendar_visible_month(driver, behavior_id, target_dt)
    time.sleep(0.3)

    for step in range(48):
        title_text = js_get_visible_calendar_title(driver, calendar_prefix)
        if title_text:
            visible_month, visible_year = _parse_calendar_title_month_year(title_text)
            if visible_month == target_month and visible_year == target_year:
                return True

            arrow_id = (
                prev_id
                if (visible_year, visible_month) > (target_year, target_month)
                else next_id
            )
            click_element_by_id(driver, arrow_id)
            time.sleep(0.25)
            continue

        if step % 4 == 3:
            set_calendar_visible_month(driver, behavior_id, target_dt)
        time.sleep(0.25)

    return False


def pick_calendar_day(driver, calendar_prefix: str, target_dt: datetime):
    """Click the calendar day cell that matches target_dt."""
    day_title = format_calendar_day_title(target_dt)
    clicked = driver.execute_script(
        """
        const prefix = arguments[0];
        const title = arguments[1];
        const cells = document.querySelectorAll('div[id*="' + prefix + '_day_"]');
        for (const cell of cells) {
            if ((cell.getAttribute('title') || '').trim() !== title) {
                continue;
            }
            const calendar = cell.closest('.ajax__calendar');
            if (calendar && calendar.offsetParent === null) {
                continue;
            }
            cell.click();
            return true;
        }
        return false;
        """,
        calendar_prefix,
        day_title,
    )
    if not clicked:
        raise RuntimeError(f"Calendar day not found for {day_title}.")
    time.sleep(0.3)


def open_date_calendar_popup(driver, input_id: str, behavior_id: str, calendar_prefix: str):
    """Open the Ajax calendar attached to a readonly date input."""
    WebDriverWait(driver, MAX_ELEMENT_WAIT_SECONDS).until(
        lambda d: bool(
            d.execute_script(
                """
                const element = document.getElementById(arguments[0]);
                return !!(element && element.offsetParent !== null);
                """,
                input_id,
            )
        )
    )

    opened = driver.execute_script(
        """
        const input = document.getElementById(arguments[0]);
        const behavior = window.$find ? window.$find(arguments[1]) : null;
        if (behavior && typeof behavior.show === 'function') {
            try {
                behavior.show();
                return true;
            } catch (e) {}
        }
        if (input) {
            input.click();
            return true;
        }
        return false;
        """,
        input_id,
        behavior_id,
    )
    if not opened:
        click_element_by_id(driver, input_id)

    WebDriverWait(driver, MAX_ELEMENT_WAIT_SECONDS).until(
        lambda d: js_is_calendar_days_visible(d, calendar_prefix)
    )


def wait_for_date_input_value(driver, input_id: str, expected_value: str):
    """Wait until a date textbox shows the expected DD/MM/YYYY value."""
    WebDriverWait(driver, MAX_ELEMENT_WAIT_SECONDS).until(
        lambda d, value=expected_value, field_id=input_id: (
            js_get_input_value_by_id(d, field_id) == value
        )
    )


def select_date_using_calendar_ui(
    driver,
    input_id: str,
    calendar_prefix: str,
    target_dt: datetime,
    ):
    """Select a date via calendar UI with CalendarBehavior fallbacks."""
    expected_value = format_dd_mm_yyyy(target_dt)
    last_error: Exception | None = None

    for attempt in range(1, MAX_CALENDAR_SELECT_ATTEMPTS + 1):
        try:
            disable_date_input_postbacks(driver)
            dismiss_active_calendar_popup(driver)
            open_date_calendar_popup(driver, input_id, calendar_prefix, calendar_prefix)

            month_ready = navigate_calendar_to_month(
                driver,
                calendar_prefix,
                calendar_prefix,
                target_dt,
            )
            if month_ready:
                try:
                    pick_calendar_day(driver, calendar_prefix, target_dt)
                except RuntimeError as day_exc:
                    print(
                        f"Calendar day click failed ({day_exc}); "
                        "using CalendarBehavior date fallback."
                    )
                    if not set_calendar_behavior_date(
                        driver, input_id, calendar_prefix, target_dt
                    ):
                        raise
            else:
                print(
                    f"Calendar month navigation failed on attempt {attempt}; "
                    "using CalendarBehavior date fallback."
                )
                if not set_calendar_behavior_date(
                    driver, input_id, calendar_prefix, target_dt
                ):
                    raise RuntimeError(
                        f"Could not navigate calendar {calendar_prefix} to "
                        f"{target_dt.strftime('%B %Y')}."
                    )

            wait_for_date_input_value(driver, input_id, expected_value)
            print(f"Date {expected_value} set for {input_id} (attempt {attempt}).")
            return
        except (StaleElementReferenceException, WebDriverException) as exc:
            if is_browser_session_lost(exc):
                raise
            last_error = exc
            print(f"Calendar selection attempt {attempt} WebDriver error: {exc}")
            dismiss_active_calendar_popup(driver)
            time.sleep(0.75)
        except Exception as exc:
            last_error = exc
            print(f"Calendar selection attempt {attempt} failed: {exc}")
            dismiss_active_calendar_popup(driver)
            if attempt >= MAX_CALENDAR_SELECT_ATTEMPTS - 1:
                if set_calendar_behavior_date(driver, input_id, calendar_prefix, target_dt):
                    try:
                        wait_for_date_input_value(driver, input_id, expected_value)
                        print(
                            f"Date {expected_value} set via behavior fallback "
                            f"(attempt {attempt})."
                        )
                        return
                    except Exception:
                        pass
            time.sleep(0.75)

    if last_error is not None:
        raise last_error
    raise RuntimeError(f"Could not select date {expected_value} in calendar UI.")


def set_date_range_and_show(driver, start_dt: datetime, end_dt: datetime):
    """
    Set From/To dates on Revenue Report and click Show.

    Calendar UI alone updates the browser DOM but not ASP.NET ViewState. The portal
    only persists dates after each field's partial postback (txtFromDate / txtToDate).
    From is synced first (which may reset To on the server), then To is set and synced.
    """
    ensure_not_502_page(driver)
    if end_dt < start_dt:
        raise ValueError("Interval end date must be on or after start date.")

    from_value = format_dd_mm_yyyy(start_dt)
    to_value = format_dd_mm_yyyy(end_dt)
    print(f"Submitting date interval: {from_value} to {to_value}")

    # Previous interval exports/postbacks can leave the UpdatePanel mid-update.
    wait_for_update_panel_idle(driver)

    select_date_using_calendar_ui(
        driver,
        FROM_DATE_INPUT_ID,
        FROM_CALENDAR_BEHAVIOR_ID,
        start_dt,
    )
    state = read_revenue_report_date_state(driver)
    assert_date_state_matches(
        state,
        from_value,
        to_value,
        "after From calendar selection",
        verify_to=False,
    )
    print(f"From date set in calendar: {state.get('fromInput', '')}")

    print("Syncing From date to server (partial postback)...")
    trigger_date_postback(driver, FROM_DATE_POSTBACK_TARGET)
    state = read_revenue_report_date_state(driver)
    assert_date_state_matches(
        state,
        from_value,
        to_value,
        "after From postback",
        verify_to=False,
    )
    print(f"From date on server: {state.get('fromInput', '')}")

    print("Setting To date after From server sync...")
    dismiss_active_calendar_popup(driver)
    driver.execute_script(
        """
        const field = document.getElementById(arguments[0]);
        if (field && field.scrollIntoView) {
            field.scrollIntoView({block: 'center', inline: 'nearest'});
        }
        """,
        TO_DATE_INPUT_ID,
    )
    time.sleep(0.3)
    select_date_using_calendar_ui(
        driver,
        TO_DATE_INPUT_ID,
        TO_CALENDAR_BEHAVIOR_ID,
        end_dt,
    )
    state = read_revenue_report_date_state(driver)
    assert_date_state_matches(state, from_value, to_value, "after To calendar selection")

    print("Syncing To date to server (partial postback)...")
    trigger_date_postback(driver, TO_DATE_POSTBACK_TARGET)
    state = read_revenue_report_date_state(driver)
    assert_date_state_matches(state, from_value, to_value, "after To postback")
    print(
        "Date values before Show: "
        f"From={state.get('fromInput', '')}, To={state.get('toInput', '')}; "
        f"calendar From={state.get('fromBehavior', '')}, "
        f"To={state.get('toBehavior', '')}"
    )

    wait_for_visible_element(driver, By.ID, "MainContentPlaceHolder_btnShow")
    click_element_by_id(driver, "MainContentPlaceHolder_btnShow")
    wait_for_show_postback(driver)
    ensure_not_502_page(driver)

    post_show_state = read_revenue_report_date_state(driver)
    print(
        "Date values after Show: "
        f"From={post_show_state.get('fromInput', '')}, "
        f"To={post_show_state.get('toInput', '')}"
    )
    assert_date_state_matches(
        post_show_state,
        from_value,
        to_value,
        "after Show postback",
    )
    reset_report_pagination_to_first_page(driver)


def wait_for_report_results(driver) -> bool:
    """Wait for the Revenue Report grid or a visible no-records message."""
    deadline = time.time() + MAX_ELEMENT_WAIT_SECONDS

    while time.time() < deadline:
        ensure_not_502_page(driver)

        try:
            grid = driver.find_element(By.ID, "MainContentPlaceHolder_grdView")
            if grid.is_displayed():
                return True
        except Exception:
            pass

        messages = collect_visible_messages(driver)
        for message in messages:
            if "no record" in message.lower():
                print(f"Report message: {message}")
                return False

        time.sleep(POLL_SECONDS)

    raise RuntimeError("Revenue Report results did not load in time.")


def get_current_report_page_number(driver) -> int:
    """Return current report page number from the pagination row."""
    try:
        current = driver.find_element(
            By.CSS_SELECTOR,
            "#MainContentPlaceHolder_grdView tr.pagination span",
        )
        return int((current.text or "").strip())
    except Exception:
        return 1


def reset_report_pagination_to_first_page(driver):
    """
    Move the Revenue Report grid back to page 1.

    After Show for a new month interval, ASP.NET can keep the previous
    pagination index (e.g. still on page 3 from the last interval).
    """
    has_pagination = driver.execute_script(
        """
        return !!document.querySelector('#MainContentPlaceHolder_grdView tr.pagination');
        """
    )
    if not has_pagination:
        return

    current_page = get_current_report_page_number(driver)
    if current_page <= 1:
        print("Report pagination is on page 1.")
        return

    print(f"Resetting report pagination from page {current_page} to page 1...")

    clicked_page_one = driver.execute_script(
        """
        const links = document.querySelectorAll('#MainContentPlaceHolder_grdView tr.pagination a');
        for (const link of links) {
            if ((link.textContent || '').trim() === '1') {
                link.click();
                return true;
            }
        }
        return false;
        """,
    )

    if clicked_page_one:
        WebDriverWait(driver, MAX_ELEMENT_WAIT_SECONDS).until(
            lambda d: get_current_report_page_number(d) == 1
        )
        ensure_not_502_page(driver)
        print("Report pagination reset to page 1.")
        return

    # Fallback: step to the lowest numbered page until we reach page 1.
    for _ in range(25):
        current_page = get_current_report_page_number(driver)
        if current_page <= 1:
            print("Report pagination reset to page 1.")
            return

        moved = driver.execute_script(
            """
            const current = parseInt(arguments[0], 10);
            const links = Array.from(
                document.querySelectorAll('#MainContentPlaceHolder_grdView tr.pagination a')
            );
            let target = null;
            let targetPage = current;
            for (const link of links) {
                const text = (link.textContent || '').trim();
                if (!/^\\d+$/.test(text)) {
                    continue;
                }
                const page = parseInt(text, 10);
                if (page < current && page < targetPage) {
                    targetPage = page;
                    target = link;
                }
            }
            if (!target) {
                return false;
            }
            target.click();
            return true;
            """,
            str(current_page),
        )
        if not moved:
            break

        previous_page = current_page
        WebDriverWait(driver, MAX_ELEMENT_WAIT_SECONDS).until(
            lambda d, prev=previous_page: get_current_report_page_number(d) < prev
        )
        time.sleep(0.3)

    final_page = get_current_report_page_number(driver)
    if final_page != 1:
        raise RuntimeError(
            f"Could not reset report pagination to page 1 (still on page {final_page})."
        )
    print("Report pagination reset to page 1.")


def open_next_report_page(driver) -> bool:
    """Open the next report page when pagination is available."""
    current_page = get_current_report_page_number(driver)
    next_link = None
    next_page = None

    try:
        links = driver.find_elements(
            By.CSS_SELECTOR,
            "#MainContentPlaceHolder_grdView tr.pagination a",
        )
    except Exception:
        return False

    for link in links:
        text = (link.text or "").strip()
        if not text.isdigit():
            continue
        page_number = int(text)
        if page_number > current_page and (next_page is None or page_number < next_page):
            next_link = link
            next_page = page_number

    if next_page is None:
        return False

    clicked = driver.execute_script(
        """
        const links = document.querySelectorAll('#MainContentPlaceHolder_grdView tr.pagination a');
        for (const link of links) {
            if ((link.textContent || '').trim() === arguments[0]) {
                link.click();
                return true;
            }
        }
        return false;
        """,
        str(next_page),
    )
    if not clicked:
        return False

    WebDriverWait(driver, MAX_ELEMENT_WAIT_SECONDS).until(
        lambda d: get_current_report_page_number(d) == next_page
    )
    ensure_not_502_page(driver)
    return True


def wait_for_new_tab(driver, existing_handles: list[str]) -> str:
    """Wait for the click to open a new browser tab and return its handle."""
    WebDriverWait(driver, MAX_ELEMENT_WAIT_SECONDS).until(
        lambda d: len(d.window_handles) > len(existing_handles)
    )
    new_handles = [handle for handle in driver.window_handles if handle not in existing_handles]
    if not new_handles:
        raise RuntimeError("New tab did not open after clicking Debit Adjustment Count.")
    return new_handles[0]


def export_current_debit_adjustment_tab(driver, plaza_name: str, settlement_date_text: str):
    """Export the report in the currently open Debit Adjustment tab and move the download."""
    ensure_not_502_page(driver)
    export_button = wait_for_visible_element(driver, By.ID, "MainContentPlaceHolder_btnExcel1")
    download_snapshot = snapshot_download_folder(DOWNLOADS_FOLDER)

    try:
        export_button.click()
    except Exception:
        driver.execute_script("arguments[0].click();", export_button)

    downloaded_file = wait_for_new_download(download_snapshot, DOWNLOADS_FOLDER)
    move_download_to_plaza_folder(downloaded_file, plaza_name, settlement_date_text)


def export_debit_adjustment_reports_for_current_interval(driver, plaza_name: str):
    """Export reports for all rows where Debit Adjustment Count is greater than 1."""
    report_loaded = wait_for_report_results(driver)
    if not report_loaded:
        print("No rows found for this interval. Moving to next interval.")
        return 0

    reset_report_pagination_to_first_page(driver)

    exported_count = 0

    while True:
        report_rows = driver.find_elements(
            By.CSS_SELECTOR,
            "#MainContentPlaceHolder_grdView tr.GvGrid",
        )

        if not report_rows:
            print("No report rows found for this interval. Moving to next interval.")
            break

        for row_number in range(len(report_rows)):
            settlement_text = ""
            debit_count = 0
            debit_link_id = ""

            for attempt in range(3):
                try:
                    report_rows = driver.find_elements(
                        By.CSS_SELECTOR,
                        "#MainContentPlaceHolder_grdView tr.GvGrid",
                    )
                    if row_number >= len(report_rows):
                        debit_count = 0
                        break
                    current_row = report_rows[row_number]
                    settlement_text = (
                        current_row.find_element(
                            By.CSS_SELECTOR,
                            "span[id*='lbl_settlementdate_']",
                        ).text.strip()
                    )
                    debit_link = current_row.find_element(
                        By.CSS_SELECTOR,
                        "a[id*='lbl_dbtadjcnt_']",
                    )
                    debit_text = debit_link.text.strip()
                    debit_link_id = (debit_link.get_attribute("id") or "").strip()
                    debit_count = int(float(debit_text))
                    break
                except StaleElementReferenceException:
                    if attempt == 2:
                        raise
                    time.sleep(0.5)
                except ValueError:
                    print(f"Skipping row with non-numeric Debit Adjustment Count: {debit_text}")
                    debit_count = 0
                    break

            if debit_count <= 1:
                continue

            print(
                f"Settlement {settlement_text}: Debit Adjustment Count is {debit_count}. Exporting report."
            )
            original_tab = driver.current_window_handle
            existing_handles = list(driver.window_handles)

            if debit_link_id:
                click_element_by_id(driver, debit_link_id)
            else:
                raise RuntimeError(
                    f"Could not resolve Debit Adjustment link id for row {row_number}."
                )

            new_tab = wait_for_new_tab(driver, existing_handles)
            driver.switch_to.window(new_tab)

            try:
                export_current_debit_adjustment_tab(driver, plaza_name, settlement_text)
                exported_count += 1
            finally:
                try:
                    driver.close()
                except Exception:
                    pass
                driver.switch_to.window(original_tab)
                ensure_not_502_page(driver)

        if not open_next_report_page(driver):
            break

    print(f"Exported {exported_count} Debit Adjustment report(s) for this interval.")
    return exported_count


def process_reference_file(
    excel_file_path: str = REFERENCE_EXCEL,
    headless: bool = False,
    date_range_start: str = DATE_RANGE_START,
    date_range_end: str = DATE_RANGE_END,
     ) -> bool:
    """Read pending rows from Excel and attempt YES Bank login for each."""
    reference_path = SCRIPT_DIR / excel_file_path
    if not reference_path.exists():
        print(f"Reference file not found: {reference_path}")
        return False

    df = pd.read_excel(reference_path, engine="openpyxl")
    try:
        col_map = resolve_required_columns(df)
    except ValueError as exc:
        print(str(exc))
        return False

    plaza_col = col_map["plaza_name"]
    username_col = col_map["username"]
    password_col = col_map["password"]
    done_col = col_map["done"]

    pending_rows = df[
        (df[done_col].isna())
        | (df[done_col].astype(str).str.strip().str.lower() != "done")
    ]

    print(f"Total rows: {len(df)}")
    print(f"Pending rows: {len(pending_rows)}")

    if pending_rows.empty:
        print("No pending rows to process.")
        return True

    try:
        intervals = split_date_range_to_month_intervals(date_range_start, date_range_end)
    except ValueError as exc:
        print(str(exc))
        return False

    print(
        f"Using {len(intervals)} month interval(s) from {date_range_start} to {date_range_end}"
    )
    for interval_start, interval_end in intervals:
        print(
            f"  Interval: {interval_start.strftime('%Y-%m-%d')} to {interval_end.strftime('%Y-%m-%d')}"
        )

    success_count = 0
    fail_count = 0

    for idx, row in pending_rows.iterrows():
        plaza_name = str(row[plaza_col]).strip() if pd.notna(row[plaza_col]) else ""
        username = str(row[username_col]).strip() if pd.notna(row[username_col]) else ""
        password = str(row[password_col]).strip() if pd.notna(row[password_col]) else ""

        print("\n" + "=" * 80)
        print(f"Processing row {idx + 2}")
        print(f"Plaza: {plaza_name}")
        print(f"Username: {username}")
        print("=" * 80)

        if not plaza_name or not username or not password:
            print("Skipping row due to missing Plaza_name/Username/Password.")
            fail_count += 1
            continue

        interval_index = 0
        row_completed = False
        row_failed = False
        browser_attempt = 0
        driver = None

        while interval_index < len(intervals) and not row_failed:
            try:
                if driver is None:
                    browser_attempt += 1
                    print(
                        f"Browser attempt {browser_attempt}/{MAX_INTERVAL_ATTEMPTS} "
                        f"for plaza {plaza_name}, interval {interval_index + 1}/{len(intervals)}"
                    )
                    driver = setup_driver(headless=headless)
                    login_ok = login_once(driver, username=username, password=password)

                    if not login_ok:
                        print("Login failed. Login form is still visible after submit.")
                        messages = collect_visible_messages(driver)
                        for message in messages:
                            print(f"Visible message: {message}")
                        fail_count += 1
                        row_failed = True
                        break

                    print("Login successful.")
                    navigate_to_revenue_report(driver)
                    print("Revenue Report page opened.")
                    select_only_client_option(driver)
                    print("Client selected.")

                interval_start, interval_end = intervals[interval_index]
                set_date_range_and_show(driver, interval_start, interval_end)
                export_debit_adjustment_reports_for_current_interval(driver, plaza_name)
                interval_index += 1
                browser_attempt = 0

            except Exception as exc:
                error_text = str(exc)
                print(f"Interval processing error: {exc}")

                if driver:
                    try:
                        driver.quit()
                    except Exception:
                        pass
                    driver = None

                if browser_attempt >= MAX_INTERVAL_ATTEMPTS:
                    fail_count += 1
                    row_failed = True
                    break

                if "502 gateway/proxy error page" in error_text:
                    print("Detected 502 page. Restarting browser and retrying the same interval.")
                elif is_browser_session_lost(exc):
                    print("Browser session lost. Restarting browser and retrying the same interval.")
                else:
                    print("Error during interval. Restarting browser and retrying the same interval.")

        if driver:
            try:
                driver.quit()
            except Exception:
                pass
            driver = None

        if interval_index >= len(intervals) and not row_failed:
            mark_done_in_excel(str(reference_path), idx, done_col)
            success_count += 1
            row_completed = True

        if not row_completed and interval_index < len(intervals):
            print(f"Failed to complete row {idx + 2}.")

    print("\n" + "=" * 80)
    print("YES BANK LOGIN SUMMARY")
    print("=" * 80)
    print(f"Successful: {success_count}")
    print(f"Failed: {fail_count}")
    print("=" * 80)

    return fail_count == 0


def main():
    process_reference_file(
        excel_file_path=REFERENCE_EXCEL,
        headless=False,
        date_range_start=DATE_RANGE_START,
        date_range_end=DATE_RANGE_END,
    )


if __name__ == "__main__":
    main()
